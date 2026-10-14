"""Render the integrated Thai Chapter 4 from the prose template and the claim registry.

Template syntax (reports/templates/chapter4_th.md)
    {{claim_id}}          replaced by the claim's display value (formatted once, from full precision)
    [[ASSET:ID]]          a figure or table insertion point; figures are embedded in the DOCX
    # / ## / ###          headings;   '> ' lines   a boxed finding (one-cell table in the DOCX)
    <!-- ... -->          editorial comments, dropped from the outputs

Unknown claim ids or asset ids stop the render with a list, so no stale or invented number
can reach the report.  The DOCX reuses the styles of the author's report (reports/source/),
whose body is replaced; the original file is never modified.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

import pandas as pd

from . import config as C
from .reporting import load_assets, load_claims

TEMPLATE = C.REPORTS / "templates" / "chapter4_th.md"
OUT_MD = C.REPORTS / "chapter4_complete_th.md"
OUT_DOCX = C.REPORTS / "chapter4_complete_th.docx"
SOURCE_DOCX = C.REPORTS / "source" / "รายงานวิจัยฉบับสมบูรณ์_LASTEST.docx"


def _numbering(text: str, assets: pd.DataFrame) -> pd.DataFrame:
    ids = re.findall(r"\[\[(?:ASSET|TABLE):([A-Z0-9_]+)\]\]", text)
    rows, nf, nt, na = [], 0, 0, 0
    appendix_at = text.find("## ภาคผนวก")
    for i in ids:
        a = assets[assets.asset_id == i]
        if a.empty:
            continue
        a = a.iloc[0]
        if a.kind == "figure":
            nf += 1
            num = f"รูปที่ {nf}"
        elif appendix_at >= 0 and text.find(i) > appendix_at:
            na += 1
            num = f"ตารางผนวกที่ ก{na}"
        else:
            nt += 1
            num = f"ตารางที่ {nt}"
        rows.append(dict(asset_id=i, kind=a.kind, number=num, caption_th=a.caption_th, file=a.file,
                         data_file=a.get("data_file", ""), section=a.section, status="placed"))
    return pd.DataFrame(rows)


def render() -> dict:
    claims = load_claims()
    assets = load_assets()
    text = TEMPLATE.read_text(encoding="utf-8")
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lookup = dict(zip(claims.claim_id, claims.display.astype(str)))
    # {{abs:id}} drops the sign where the Thai verb already carries the direction ("ลดลง ...")
    for cid in set(re.findall(r"\{\{abs:([^}]+)\}\}", text)):
        if cid in lookup:
            lookup["abs:" + cid] = lookup[cid].lstrip("-+−")
    missing = sorted(set(re.findall(r"\{\{([^}]+)\}\}", text)) - set(lookup))
    bad_assets = sorted(set(re.findall(r"\[\[(?:ASSET|TABLE):([A-Z0-9_]+)\]\]", text)) - set(assets.asset_id))
    if missing or bad_assets:
        raise KeyError(f"unknown claim ids {missing}; unknown asset ids {bad_assets}")
    used = sorted(set(re.findall(r"\{\{([^}]+)\}\}", text)))
    body = re.sub(r"\{\{([^}]+)\}\}", lambda m: lookup[m.group(1)], text)
    num = _numbering(body, assets)
    nmap = num.set_index("asset_id") if len(num) else pd.DataFrame()

    def marker(m):
        a = nmap.loc[m.group(1)]
        kind = "รูป" if a.kind == "figure" else "ตาราง"
        return f"> [ตำแหน่ง{kind} {m.group(1)} — {a.number} {a.caption_th}; ไฟล์: `{a.file}`]"
    def md_table(m):
        a = nmap.loc[m.group(1)]
        t = pd.read_csv(C.ROOT / a.file, encoding="utf-8-sig", dtype=str).fillna("")
        head = "| " + " | ".join(t.columns) + " |\n|" + "---|" * len(t.columns)
        rows = "\n".join("| " + " | ".join(r) + " |" for r in t.values.tolist())
        return f"{a.number} {a.caption_th}\n\n{head}\n{rows}"
    md = re.sub(r"\[\[TABLE:([A-Z0-9_]+)\]\]", md_table, body)
    md = re.sub(r"\[\[ASSET:([A-Z0-9_]+)\]\]", marker, md)
    OUT_MD.write_text(md, encoding="utf-8")
    num.to_csv(C.REPORTS / "asset_numbering_map.csv", index=False, encoding="utf-8-sig")
    reg = claims[claims.claim_id.isin([u.removeprefix("abs:") for u in used])].copy()
    reg["report_location"] = "reports/chapter4_complete_th.md"
    reg.to_csv(C.REPORTS / "claims_used_in_report.csv", index=False, encoding="utf-8-sig")
    _docx(body, nmap)
    print(f"report rendered: {len(used)} claims, {len(num)} assets -> {OUT_MD.relative_to(C.ROOT)}, {OUT_DOCX.relative_to(C.ROOT)}")
    return dict(claims_used=len(used), assets=len(num))


def _docx(body: str, nmap: pd.DataFrame):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    d = docx.Document(str(SOURCE_DOCX))
    bodyel = d.element.body
    sect = bodyel[-1] if bodyel[-1].tag.endswith("sectPr") else None
    for el in list(bodyel):
        if el is not sect:
            bodyel.remove(el)
    styles = {s.name for s in d.styles}

    def st(name, fallback="Normal"):
        return name if name in styles else fallback

    def add_runs(par, text):
        # **bold** and `code` minimal inline markup
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if part.startswith("**") and part.endswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`"):
                par.add_run(part[1:-1])
            elif part:
                par.add_run(part)

    box: list[str] = []

    def flush_box():
        if not box:
            return
        t = d.add_table(rows=1, cols=1)
        t.style = d.styles["Table Grid"] if "Table Grid" in styles else t.style
        cell = t.cell(0, 0)
        cell.paragraphs[0].text = ""
        for i, line in enumerate(box):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            add_runs(p, line)
        d.add_paragraph()
        box.clear()

    for raw in body.split("\n"):
        line = raw.rstrip()
        tm = re.match(r"\[\[TABLE:([A-Z0-9_]+)\]\]", line.strip())
        if tm:
            flush_box()
            a = nmap.loc[tm.group(1)]
            d.add_paragraph(f"{a.number} {a.caption_th}", style=st("Caption"))
            t = pd.read_csv(C.ROOT / a.file, encoding="utf-8-sig", dtype=str).fillna("")
            tab = d.add_table(rows=len(t) + 1, cols=len(t.columns))
            if "Table Grid" in styles:
                tab.style = d.styles["Table Grid"]
            for j, col in enumerate(t.columns):
                cell = tab.cell(0, j); cell.text = col
                for r_ in cell.paragraphs[0].runs:
                    r_.bold = True
            for i, row in enumerate(t.values.tolist(), start=1):
                for j, v in enumerate(row):
                    tab.cell(i, j).text = v
                    if j >= 2:
                        tab.cell(i, j).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            d.add_paragraph()
            continue
        m = re.match(r"\[\[ASSET:([A-Z0-9_]+)\]\]", line.strip())
        if line.startswith("> ") and not m:
            box.append(line[2:])
            continue
        flush_box()
        if not line.strip():
            continue
        if m:
            a = nmap.loc[m.group(1)]
            if a.kind == "figure":
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(C.ROOT / a.file), width=Cm(15.5))
                cap = d.add_paragraph(f"{a.number} {a.caption_th}", style=st("Caption"))
            else:
                cap = d.add_paragraph(f"{a.number} {a.caption_th}", style=st("Caption"))
                d.add_paragraph(f"[ตำแหน่งตาราง {m.group(1)}: ข้อมูลเต็มอยู่ที่ {a.file}]")
            continue
        h = re.match(r"^(#{1,3})\s+(.*)$", line)
        if h:
            lvl = len(h.group(1))
            d.add_paragraph(h.group(2), style=st(f"Heading {lvl}"))
            continue
        if line.startswith("- "):
            p = d.add_paragraph(style=st("List Bullet"))
            add_runs(p, line[2:])
            continue
        p = d.add_paragraph(style="Normal")
        add_runs(p, line)
    flush_box()
    d.save(str(OUT_DOCX))
