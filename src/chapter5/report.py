"""Render the Thai Chapter 5 from reports/templates/chapter5_th.md and the claim registry.

Template syntax (same as Chapter 4):
    {{claim_id}}      the claim's display value;  {{abs:claim_id}} drops the sign
    [[ASSET:ID]]      figure/table insertion marker (figures embedded in the DOCX)
    [[TABLE:ID]]      the table printed in full (short tables only)
    <!-- ... -->      editorial comments, dropped

Unknown claim or asset ids stop the render, so no stale or invented number reaches the text.
"""
from __future__ import annotations

import re

import pandas as pd

from chapter4 import config as C4
from .reporting import load_registry

TEMPLATE = C4.REPORTS / "templates" / "chapter5_th.md"
SOURCE_DOCX = C4.REPORTS / "source" / "รายงานวิจัยฉบับสมบูรณ์_LASTEST.docx"


def _paths(S):
    if S.name == "canonical":
        return C4.REPORTS / "chapter5_complete_th.md", C4.REPORTS / "chapter5_complete_th.docx", C4.REPORTS / "chapter5_asset_numbering_map.csv"
    d = S.out / "report"
    d.mkdir(parents=True, exist_ok=True)
    return d / "chapter5_complete_th.md", d / "chapter5_complete_th.docx", d / "chapter5_asset_numbering_map.csv"


def _numbering(text, assets):
    ids = re.findall(r"\[\[(?:ASSET|TABLE):([A-Z0-9_]+)\]\]", text)
    rows, nf, nt, na = [], 0, 0, 0
    app = text.find("## ภาคผนวก")
    seen = set()
    for i in ids:
        if i in seen:
            continue
        seen.add(i)
        a = assets[assets.asset_id == i].iloc[0]
        if a.kind == "figure":
            nf += 1
            num = f"รูปที่ 5.{nf}"
        elif app >= 0 and text.find(i) > app:
            na += 1
            num = f"ตารางผนวกที่ 5ก.{na}"
        else:
            nt += 1
            num = f"ตารางที่ 5.{nt}"
        rows.append(dict(asset_id=i, kind=a.kind, number=num, caption_th=a.caption_th, file=a.file,
                         data_file=a.get("data_file", ""), section=a.section, status="placed"))
    return pd.DataFrame(rows)


def render(S) -> dict:
    claims = load_registry(S, "claims")
    assets = load_registry(S, "assets")
    text = TEMPLATE.read_text(encoding="utf-8")
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lookup = dict(zip(claims.claim_id, claims.display.astype(str)))
    for cid in set(re.findall(r"\{\{abs:([^}]+)\}\}", text)):
        if cid in lookup:
            lookup["abs:" + cid] = lookup[cid].lstrip("-+−")
    missing = sorted(set(re.findall(r"\{\{([^}]+)\}\}", text)) - set(lookup))
    bad = sorted(set(re.findall(r"\[\[(?:ASSET|TABLE):([A-Z0-9_]+)\]\]", text)) - set(assets.asset_id))
    if missing or bad:
        raise KeyError(f"unknown claim ids {missing}; unknown asset ids {bad}")
    used = sorted(set(re.findall(r"\{\{([^}]+)\}\}", text)))
    body = re.sub(r"\{\{([^}]+)\}\}", lambda m: lookup[m.group(1)], text)
    num = _numbering(body, assets)
    nmap = num.set_index("asset_id")
    out_md, out_docx, out_map = _paths(S)

    def marker(m):
        a = nmap.loc[m.group(1)]
        kind = "รูป" if a.kind == "figure" else "ตาราง"
        return f"> [ตำแหน่ง{kind} {m.group(1)} — {a.number} {a.caption_th}; ไฟล์: `{a.file}`]"

    def md_table(m):
        a = nmap.loc[m.group(1)]
        t = pd.read_csv(C4.ROOT / a.file, encoding="utf-8-sig", dtype=str).fillna("")
        head = "| " + " | ".join(t.columns) + " |\n|" + "---|" * len(t.columns)
        rows = "\n".join("| " + " | ".join(r) + " |" for r in t.values.tolist())
        return f"**{a.number} {a.caption_th}** (ไฟล์: `{a.file}`)\n\n{head}\n{rows}"
    md = re.sub(r"\[\[TABLE:([A-Z0-9_]+)\]\]", md_table, body)
    md = re.sub(r"\[\[ASSET:([A-Z0-9_]+)\]\]", marker, md)
    out_md.write_text(md, encoding="utf-8")
    num.to_csv(out_map, index=False, encoding="utf-8-sig")
    reg = claims[claims.claim_id.isin([u.removeprefix("abs:") for u in used])].copy()
    reg["report_location"] = str(out_md.relative_to(C4.ROOT)) if out_md.is_relative_to(C4.ROOT) else str(out_md)
    reg.to_csv(out_map.parent / "chapter5_claims_used_in_report.csv", index=False, encoding="utf-8-sig")
    _docx(body, nmap, out_docx)
    print(f"report: {len(used)} claims, {len(num)} assets -> {out_md}")
    return dict(claims=len(used), assets=len(num), md=out_md, docx=out_docx)


def _docx(body, nmap, out):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
    d = docx.Document(str(SOURCE_DOCX))
    el = d.element.body
    sect = el[-1] if el[-1].tag.endswith("sectPr") else None
    for e in list(el):
        if e is not sect:
            el.remove(e)
    styles = {s.name for s in d.styles}
    st = lambda n, f="Normal": n if n in styles else f

    def runs(p, text):
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\$[^$]+\$)", text):
            if part.startswith("**") and part.endswith("**"):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`"):
                p.add_run(part[1:-1])
            elif part.startswith("$") and part.endswith("$") and len(part) > 2:
                r = p.add_run(part[1:-1]); r.italic = True
            elif part:
                p.add_run(part)

    box = []

    def flush():
        if not box:
            return
        t = d.add_table(rows=1, cols=1)
        if "Table Grid" in styles:
            t.style = d.styles["Table Grid"]
        c = t.cell(0, 0); c.paragraphs[0].text = ""
        for i, line in enumerate(box):
            runs(c.paragraphs[0] if i == 0 else c.add_paragraph(), line)
        d.add_paragraph()
        box.clear()

    def table(a):
        d.add_paragraph(f"{a.number} {a.caption_th}", style=st("Caption"))
        t = pd.read_csv(C4.ROOT / a.file, encoding="utf-8-sig", dtype=str).fillna("")
        tab = d.add_table(rows=len(t) + 1, cols=len(t.columns))
        if "Table Grid" in styles:
            tab.style = d.styles["Table Grid"]
        for j, col in enumerate(t.columns):
            c = tab.cell(0, j); c.text = col
            for r_ in c.paragraphs[0].runs:
                r_.bold = True; r_.font.size = Pt(8)
        for i, row in enumerate(t.values.tolist(), start=1):
            for j, v in enumerate(row):
                c = tab.cell(i, j); c.text = v
                for r_ in c.paragraphs[0].runs:
                    r_.font.size = Pt(8)
        d.add_paragraph()

    in_math = False
    for raw in body.split("\n"):
        line = raw.rstrip()
        if line.strip() == "$$":
            in_math = not in_math
            continue
        if in_math:
            p = d.add_paragraph(style="Normal"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line.strip()); r.italic = True
            continue
        tm = re.match(r"\[\[TABLE:([A-Z0-9_]+)\]\]", line.strip())
        if tm:
            flush(); table(nmap.loc[tm.group(1)]); continue
        m = re.match(r"\[\[ASSET:([A-Z0-9_]+)\]\]", line.strip())
        if line.startswith("> ") and not m:
            box.append(line[2:]); continue
        flush()
        if not line.strip():
            continue
        if m:
            a = nmap.loc[m.group(1)]
            if a.kind == "figure":
                p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(C4.ROOT / a.file), width=Cm(15.5))
                d.add_paragraph(f"{a.number} {a.caption_th}", style=st("Caption"))
            else:
                t = pd.read_csv(C4.ROOT / a.file, encoding="utf-8-sig", dtype=str)
                if len(t) <= 40 and len(t.columns) <= 14:
                    table(a)
                else:
                    d.add_paragraph(f"{a.number} {a.caption_th}", style=st("Caption"))
                    d.add_paragraph(f"[ตารางเต็มอยู่ในไฟล์ {a.file}]")
            continue
        if line.startswith("|"):
            if set(line.replace("|", "").strip()) <= set("-: "):
                continue
            p = d.add_paragraph(style="Normal"); runs(p, "  ".join(c.strip() for c in line.strip("|").split("|")))
            continue
        h = re.match(r"^(#{1,3})\s+(.*)$", line)
        if h:
            d.add_paragraph(h.group(2), style=st(f"Heading {len(h.group(1))}")); continue
        if line.startswith("- "):
            p = d.add_paragraph(style=st("List Bullet")); runs(p, line[2:]); continue
        p = d.add_paragraph(style="Normal"); runs(p, line)
    flush()
    d.save(str(out))
